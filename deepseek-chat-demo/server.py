import os
import time
from collections import defaultdict
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, request, jsonify, send_from_directory
from werkzeug.utils import secure_filename

from client import send_message


ROOT = Path(__file__).parent
load_dotenv(ROOT / ".env")

app = Flask(__name__)
app.config["UPLOAD_FOLDER"] = str(ROOT / "uploads")
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

Path(app.config["UPLOAD_FOLDER"]).mkdir(parents=True, exist_ok=True)


_conversations: dict[str, list[dict]] = defaultdict(list)


def append_message(uid: str, role: str, content: str) -> None:
    _conversations[uid].append({
        "role": role,
        "content": content,
        "ts": time.time(),
    })


def get_history(uid: str, last_n: int = 50) -> list[dict]:
    return _conversations[uid][-last_n:]


def clear_history(uid: str) -> None:
    _conversations.pop(uid, None)


ALLOWED_EXTENSIONS = {
    "txt", "md", "pdf", "doc", "docx", "xls", "xlsx",
    "png", "jpg", "jpeg", "gif", "webp",
}


def allowed_file(filename: str) -> bool:
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def _read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        chunks = []
        for page in reader.pages[:20]:
            chunks.append(page.extract_text() or "")
        return "\n".join(chunks)
    except Exception as exc:  # pragma: no cover - best effort parsing
        return f"[PDF parse failed: {exc}]"


def _read_docx(path: Path) -> str:
    try:
        from docx import Document

        doc = Document(str(path))
        return "\n".join(paragraph.text for paragraph in doc.paragraphs)
    except Exception as exc:  # pragma: no cover - best effort parsing
        return f"[DOCX parse failed: {exc}]"


def _read_xlsx(path: Path) -> str:
    try:
        import openpyxl

        workbook = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
        output = []
        for sheet in workbook.worksheets[:5]:
            output.append(f"# Worksheet {sheet.title}")
            for idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=100, values_only=True), start=1):
                values = ["" if cell is None else str(cell) for cell in row]
                output.append("\t".join(values))
                if idx >= 100:
                    break
        return "\n".join(output)
    except Exception as exc:  # pragma: no cover - best effort parsing
        return f"[XLSX parse failed: {exc}]"


def extract_text(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    if ext in {"txt", "md"}:
        return _read_txt(path)
    if ext == "pdf":
        return _read_pdf(path)
    if ext in {"doc", "docx"}:
        return _read_docx(path)
    if ext in {"xls", "xlsx"}:
        return _read_xlsx(path)
    if ext in {"png", "jpg", "jpeg", "gif", "webp"}:
        return f"[Image file: {path.name}]"
    return f"[Unsupported file: {path.name}]"


def shorten(content: str, limit: int = 8000) -> str:
    content = content.strip()
    if len(content) <= limit:
        return content
    return f"{content[:limit]}\n...[truncated, total {len(content)} chars]"


@app.get("/")
def index():
    return send_from_directory(ROOT, "chat.html")


@app.post("/chat")
def chat():
    data = request.get_json(silent=True) or {}
    uid = data.get("user_id", "u_default")
    text = (data.get("message") or "").strip()
    if not text:
        return jsonify({"error": "no message"}), 400

    append_message(uid, "user", text)
    history = [{"role": m["role"], "content": m["content"]} for m in get_history(uid)]

    try:
        result = send_message(uid, history)
        reply = (
            result.get("choices", [{}])[0].get("message", {}).get("content")
            or result.get("reply")
            or result.get("response")
            or str(result)
        )
    except Exception as exc:  # pragma: no cover - forward error
        return jsonify({"error": str(exc)}), 500

    append_message(uid, "assistant", reply)
    return jsonify({"reply": reply})


@app.post("/clear")
def clear():
    uid = (request.get_json(silent=True) or {}).get("user_id", "u_default")
    clear_history(uid)
    return jsonify({"ok": True})


@app.get("/history")
def history():
    uid = request.args.get("user_id", "u_default")
    last_n = min(int(request.args.get("n", "50")), 200)
    items = [{
        "role": m["role"],
        "content": m["content"],
        "ts": m["ts"],
    } for m in get_history(uid, last_n)]
    return jsonify({"messages": items})


@app.post("/upload")
def upload():
    uid = request.form.get("user_id", "u_default")
    user_note = (request.form.get("message") or "").strip()

    files = [f for f in request.files.getlist("files") if f and f.filename]
    if not files:
        return jsonify({"error": "no files"}), 400

    saved_names = []
    abstracts = []

    for file_obj in files:
        if not allowed_file(file_obj.filename):
            return jsonify({"error": f"unsupported file type: {file_obj.filename}"}), 400

        filename = secure_filename(file_obj.filename)
        target = Path(app.config["UPLOAD_FOLDER"]) / filename
        file_obj.save(str(target))
        saved_names.append(filename)

        abstract = extract_text(target)
        abstracts.append(shorten(abstract, 6000))

    summary = "The user uploaded the following files. Review them before answering:\n\n" + "\n\n".join(abstracts)
    if user_note:
        summary += f"\n\nAdditional note: {user_note}"

    append_message(uid, "user", summary)
    history = [{"role": m["role"], "content": m["content"]} for m in get_history(uid)]

    try:
        result = send_message(uid, history)
        reply = (
            result.get("choices", [{}])[0].get("message", {}).get("content")
            or result.get("reply")
            or result.get("response")
            or str(result)
        )
    except Exception as exc:  # pragma: no cover - forward error
        return jsonify({"error": str(exc)}), 500

    append_message(uid, "assistant", reply)
    return jsonify({"reply": reply, "files": saved_names})


@app.get("/uploads/<path:filename>")
def serve_upload(filename: str):
    return send_from_directory(app.config["UPLOAD_FOLDER"], filename, as_attachment=False)


if __name__ == "__main__":
    port = int(os.getenv("PORT", "7860"))
    app.run(host="0.0.0.0", port=port, debug=True)
